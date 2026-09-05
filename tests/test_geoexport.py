"""Tests for the geospatial exporters."""
import io
import json
import zipfile

import shapefile

import pytest

from geocoord.geoexport import sanitize_filename, to_geojson, to_kml, to_shapefile_zip

FEATURES = [
    (-8.0, 39.0, {"name": "Lisboa", "value": 1}),
    (-8.61, 41.16, {"name": "Porto", "value": 2}),
]
FIELDS = ["name", "value"]


def test_geojson_structure():
    fc = json.loads(to_geojson(FEATURES).decode("utf-8"))
    assert fc["type"] == "FeatureCollection"
    assert len(fc["features"]) == 2
    f0 = fc["features"][0]
    assert f0["geometry"]["coordinates"] == [-8.0, 39.0]
    assert f0["properties"]["name"] == "Lisboa"


def test_kml_content():
    kml = to_kml(FEATURES, name_key="name").decode("utf-8")
    assert "<kml" in kml
    assert "-8.0,39.0,0" in kml
    assert "<name>Lisboa</name>" in kml


def test_shapefile_zip_roundtrip():
    data = to_shapefile_zip(FEATURES, FIELDS)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
        assert {"coordinates.shp", "coordinates.shx", "coordinates.dbf",
                "coordinates.prj"} <= names
        shp = io.BytesIO(z.read("coordinates.shp"))
        shx = io.BytesIO(z.read("coordinates.shx"))
        dbf = io.BytesIO(z.read("coordinates.dbf"))
        prj = z.read("coordinates.prj").decode("utf-8")

    assert "WGS_1984" in prj
    reader = shapefile.Reader(shp=shp, shx=shx, dbf=dbf)
    assert reader.numRecords == 2
    pts = [s.points[0] for s in reader.shapes()]
    assert pts[0][0] == -8.0 and pts[0][1] == 39.0
    recs = reader.records()
    assert recs[0]["name"] == "Lisboa"


@pytest.mark.parametrize("raw,expected", [
    ("dados das áreas de ferro.csv", "dados_das_areas_de_ferro"),
    ("sites.csv", "sites"),
    ("Coordenadas (Final).xlsx", "Coordenadas_Final"),
    ("relatório-2024.geojson", "relatorio-2024"),
    ("C:/tmp/São Tomé.csv", "Sao_Tome"),
    ("amostras_ção.shp", "amostras_cao"),
])
def test_sanitize_filename(raw, expected):
    assert sanitize_filename(raw) == expected


@pytest.mark.parametrize("raw", ["", "   ", "***", ".csv", "///"])
def test_sanitize_filename_falls_back_to_default(raw):
    assert sanitize_filename(raw) == "converted"
    assert sanitize_filename(raw, default="layer") == "layer"


def test_sanitize_filename_is_idempotent():
    once = sanitize_filename("dados das áreas (v2).csv")
    assert sanitize_filename(once) == once


def test_sanitize_filename_trims_length():
    assert len(sanitize_filename("a" * 200)) == 60


def test_shapefile_zip_uses_base_name_for_internal_layer():
    data = to_shapefile_zip(FEATURES, FIELDS, base_name="dados das áreas.csv")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
    assert {"dados_das_areas.shp", "dados_das_areas.shx",
            "dados_das_areas.dbf", "dados_das_areas.prj"} <= names


def test_shapefile_long_field_names_are_truncated():
    feats = [(-8.0, 39.0, {"a_very_long_attribute_name": "x"})]
    data = to_shapefile_zip(feats, ["a_very_long_attribute_name"])
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        dbf = io.BytesIO(z.read("coordinates.dbf"))
        shp = io.BytesIO(z.read("coordinates.shp"))
        shx = io.BytesIO(z.read("coordinates.shx"))
    reader = shapefile.Reader(shp=shp, shx=shx, dbf=dbf)
    field_names = [f[0] for f in reader.fields if f[0] != "DeletionFlag"]
    assert all(len(name) <= 10 for name in field_names)


# ---------------------------------------------------------------------------
# The Excel export, which the format itself makes hazardous
# ---------------------------------------------------------------------------
def _excel_cells(frame):
    """Write the frame and read back each cell's value and openpyxl type."""
    import io

    import openpyxl

    import app

    book = openpyxl.load_workbook(io.BytesIO(app.to_excel_bytes(frame)))
    sheet = book.active
    return [(c.value, c.data_type) for row in sheet.iter_rows(min_row=2) for c in row]


def test_excel_export_does_not_write_live_formulas():
    """A converted file is usually somebody else's data, opened on your machine.

    openpyxl writes a cell whose text begins with "=" as a formula, so
    =HYPERLINK(...) in a name column arrived live and ran. The text is kept
    exactly as it was and marked as a string: nothing is altered, nothing
    executes. The CSV export has been guarded by csv_safe for a while; this was
    the same hole in the other one.
    """
    pytest.importorskip("openpyxl")
    pd = pytest.importorskip("pandas")
    frame = pd.DataFrame({"nome": ['=1+1', '=HYPERLINK("http://x","go")', '@SUM(A1)'],
                          "lat": ["38.5"] * 3})
    cells = _excel_cells(frame)
    assert all(kind != "f" for _, kind in cells), "a cell was written as a formula"
    # ...and the value is untouched, so the data still says what it said.
    assert cells[0][0] == "=1+1"


def test_excel_export_survives_a_control_character():
    """One stray byte in a notes column used to mean no Excel download at all.

    openpyxl refuses a workbook containing a C0 control character, and the
    exception took the whole download step with it.
    """
    pytest.importorskip("openpyxl")
    pd = pytest.importorskip("pandas")
    for ch in ("\x07", "\x0b", "\x00", "\x1b"):
        frame = pd.DataFrame({"nota": [f"a{ch}b"], "lat": ["38.5"]})
        cells = _excel_cells(frame)
        assert cells[0][0] == "ab", f"the control character {ch!r} was not dropped"


def test_excel_export_keeps_tab_and_newline():
    """Only the characters the format actually refuses are dropped."""
    pytest.importorskip("openpyxl")
    pd = pytest.importorskip("pandas")
    frame = pd.DataFrame({"nota": ["a\tb\nc"], "lat": ["38.5"]})
    assert _excel_cells(frame)[0][0] == "a\tb\nc"
